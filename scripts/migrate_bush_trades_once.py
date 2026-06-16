"""One-off migration: BUSH Hedge Fund trading log -> input_trades.xlsx.

Reads Week 1 and Week 2 from the BUSH workbook, dedupes carry-over rows,
maps broker instruments to Yahoo tickers, and writes the Tracker ledger format.

Not wired into DVC/CLI — run manually:
    python scripts/migrate_bush_trades_once.py
    aqre tracker ingest --force
"""

from __future__ import annotations

import re
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = Path(
    r"D:\Romain\Projects\Finance\02_Risk Analysis\BUSH Hedge Fund_trading log_SGD.xlsx"
)
OUT_PATH = PROJECT_ROOT / "data" / "input" / "portfolio" / "input_trades.xlsx"

HEADERS = [
    "trade_id",
    "ticker",
    "trade_date",
    "side",
    "action",
    "quantity",
    "price",
    "fees",
    "lot_id",
    "rolled_from_lot_id",
    "notes",
]

# AAPL short W2 #12: stopped out Jun 5 per user confirmation (use stop-loss column)
AAPL_SHORT_INFERRED_CLOSE = {"close_date": date(2026, 6, 5)}

EXPECTED_OPEN = [
    ("CVX", "long", 543.0, 184.30),
    ("INTC", "short", 40.0, 110.85),
    ("NVDA", "short", 44.0, 213.82),
    ("^GSPC", "short", 0.5, 7582.91),
]


def parse_ticker(instrument: str) -> str:
    text = str(instrument or "").strip()
    m = re.search(r"\(([A-Z0-9^=\.]+)\)", text)
    if m:
        sym = m.group(1)
        if sym == "NAS100":
            return "^NDX"
        return sym
    low = text.lower()
    if "us 500" in low or "s&p500" in low or "s&p 500" in low:
        return "^GSPC"
    if "nas100" in low or "tech 100" in low:
        return "^NDX"
    raise ValueError(f"Cannot map instrument to ticker: {instrument!r}")


def parse_direction(value: Any) -> str:
    text = str(value or "").strip().lower()
    if text in {"long", "l"}:
        return "long"
    if text in {"short", "s"}:
        return "short"
    raise ValueError(f"Invalid direction: {value!r}")


def parse_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%d-%b-%y", "%b-%d", "%d %b %Y"):
        try:
            parsed = datetime.strptime(text, fmt)
            if fmt == "%b-%d":
                parsed = parsed.replace(year=2026)
            return parsed.date()
        except ValueError:
            continue
    raise ValueError(f"Cannot parse date: {value!r}")


def parse_float(value: Any) -> float:
    if value is None:
        raise ValueError("Missing numeric value")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace(",", "").strip()
    return float(text)


def find_header_row(rows: list[tuple[Any, ...]]) -> int:
    for i, row in enumerate(rows):
        if len(row) > 2 and row[2] is not None and "Date" in str(row[2]) and "Opened" in str(row[2]):
            return i
    raise ValueError("Trade header row not found")


def parse_week_sheet(ws, week_num: int) -> list[dict[str, Any]]:
    rows = list(ws.iter_rows(values_only=True))
    hdr_idx = find_header_row(rows)
    trades: list[dict[str, Any]] = []

    for row in rows[hdr_idx + 1 :]:
        if not row or row[2] is None:
            continue
        try:
            opened = parse_date(row[2])
        except ValueError:
            continue
        if opened is None:
            continue

        trade_num = row[1]
        instrument = row[3]
        if not instrument:
            continue

        ticker = parse_ticker(instrument)
        side = parse_direction(row[4])
        qty_raw = parse_float(row[5])
        quantity = abs(qty_raw)
        fill_price = parse_float(row[6])
        stop_loss = parse_float(row[8]) if len(row) > 8 and row[8] is not None else None
        close_date = parse_date(row[11]) if len(row) > 11 else None
        exit_price = parse_float(row[12]) if len(row) > 12 and row[12] is not None else None
        notes = str(row[15]).strip() if len(row) > 15 and row[15] is not None else ""
        status = str(row[16]).strip().lower() if len(row) > 16 and row[16] is not None else ""

        lot_id = f"W{week_num}-T{trade_num}-{ticker}"
        lot_key = (ticker, opened, side, quantity, fill_price)

        trades.append(
            {
                "week": week_num,
                "trade_num": trade_num,
                "lot_id": lot_id,
                "lot_key": lot_key,
                "ticker": ticker,
                "side": side,
                "quantity": quantity,
                "open_date": opened,
                "open_price": fill_price,
                "close_date": close_date,
                "close_price": exit_price,
                "stop_loss": stop_loss,
                "status": status,
                "notes": notes,
            }
        )
    return trades


def lot_to_rows(lot: dict[str, Any]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    open_action = "buy" if lot["side"] == "long" else "sell"
    rows.append(
        {
            "trade_id": str(uuid.uuid4()),
            "ticker": lot["ticker"],
            "trade_date": lot["open_date"],
            "side": lot["side"],
            "action": open_action,
            "quantity": lot["quantity"],
            "price": lot["open_price"],
            "fees": 0.0,
            "lot_id": lot["lot_id"],
            "rolled_from_lot_id": "",
            "notes": lot["notes"],
        }
    )

    close_date = lot.get("close_date")
    close_price = lot.get("close_price")
    if close_date is not None and close_price is not None:
        close_action = "sell" if lot["side"] == "long" else "buy"
        rows.append(
            {
                "trade_id": str(uuid.uuid4()),
                "ticker": lot["ticker"],
                "trade_date": close_date,
                "side": lot["side"],
                "action": close_action,
                "quantity": lot["quantity"],
                "price": close_price,
                "fees": 0.0,
                "lot_id": lot["lot_id"],
                "rolled_from_lot_id": "",
                "notes": f"Close: {lot['notes']}" if lot["notes"] else "Close",
            }
        )
    return rows


def lot_match_key(lot: dict[str, Any]) -> tuple:
    """Identity for the same economic position leg (ignores open_date carry-over drift)."""
    return (lot["ticker"], lot["side"], lot["quantity"], lot["open_price"])


def merge_weeks(week1: list[dict], week2: list[dict]) -> list[dict]:
    """Dedupe carry-over rows; Week 2 may re-list Week 1 opens with a different date."""
    lots_by_key: dict[tuple, dict] = {}

    for lot in week1:
        lots_by_key[lot_match_key(lot)] = dict(lot)

    for lot in week2:
        key = lot_match_key(lot)
        if key not in lots_by_key:
            lots_by_key[key] = dict(lot)
            continue

        existing = lots_by_key[key]
        if lot.get("close_date"):
            existing["close_date"] = lot["close_date"]
            existing["close_price"] = lot["close_price"]
        if lot.get("status"):
            existing["status"] = lot["status"]
        if lot.get("notes") and not existing.get("notes"):
            existing["notes"] = lot["notes"]

    lots = lots_by_key

    # AAPL short inferred stop-out (sheet status Open but notes say stopped Jun 5)
    aapl_key = ("AAPL", "short", 60.0, 311.87)
    for key, lot in lots.items():
        if key == aapl_key and lot.get("close_date") is None:
            lot["close_date"] = AAPL_SHORT_INFERRED_CLOSE["close_date"]
            stop = lot.get("stop_loss")
            lot["close_price"] = stop if stop is not None else lot["open_price"]
            lot["notes"] = (lot.get("notes") or "") + " [inferred stop-out Jun 5]"

    return list(lots.values())


def build_ledger(source: Path) -> list[dict[str, Any]]:
    wb = load_workbook(source, read_only=True, data_only=True)
    week1 = parse_week_sheet(wb["Week 1"], 1)
    week2 = parse_week_sheet(wb["Week 2"], 2)
    wb.close()

    lots = merge_weeks(week1, week2)
    ledger: list[dict[str, Any]] = []
    for lot in sorted(lots, key=lambda x: (x["open_date"], x["lot_id"])):
        ledger.extend(lot_to_rows(lot))
    ledger.sort(key=lambda r: (r["trade_date"], r["lot_id"], r["action"]))
    return ledger


def open_positions_from_ledger(ledger: list[dict]) -> list[tuple[str, str, float, float]]:
    """Return (ticker, side, qty, avg_cost) for open lots."""
    from collections import defaultdict

    pos: dict[str, dict] = defaultdict(lambda: {"qty": 0.0, "side": None, "cost": 0.0})

    for row in ledger:
        lot = row["lot_id"]
        ticker = row["ticker"]
        side = row["side"]
        qty = row["quantity"]
        price = row["price"]
        action = row["action"]

        key = lot
        if pos[key]["side"] is None:
            pos[key]["side"] = side

        delta = qty if action == "buy" else -qty
        if side == "short":
            delta = -qty if action == "sell" else qty

        if side == "long":
            if action == "buy":
                pos[key]["qty"] += qty
                pos[key]["cost"] = price
            else:
                pos[key]["qty"] -= qty
        else:
            if action == "sell":
                pos[key]["qty"] += qty
                pos[key]["cost"] = price
            else:
                pos[key]["qty"] -= qty

    open_lots = []
    for lot_id, p in pos.items():
        if abs(p["qty"]) > 1e-9:
            ticker = lot_id.split("-")[-1]
            open_lots.append((ticker, p["side"], p["qty"], p["cost"]))
    return sorted(open_lots, key=lambda x: x[0])


def write_excel(ledger: list[dict], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "trades"
    ws.append(HEADERS)
    for row in ledger:
        ws.append([row[h] for h in HEADERS])
    wb.save(out_path)


def main() -> None:
    source = DEFAULT_SOURCE
    if not source.is_file():
        raise FileNotFoundError(f"BUSH workbook not found: {source}")

    ledger = build_ledger(source)
    write_excel(ledger, OUT_PATH)

    print(f"Wrote {len(ledger)} trade rows -> {OUT_PATH}")
    print("\nOpen positions (from ledger):")
    for ticker, side, qty, cost in open_positions_from_ledger(ledger):
        print(f"  {ticker:6} {side:5} qty={qty:g} entry={cost}")

    print("\nExpected open positions:")
    for ticker, side, qty, entry in EXPECTED_OPEN:
        print(f"  {ticker:6} {side:5} qty={qty:g} entry={entry}")

    gs_rows = [r for r in ledger if r["ticker"] == "GS"]
    if gs_rows:
        print("\nNote: GS prices in source look high — verify against Saxo if marks seem off:")
        for r in gs_rows:
            print(f"  {r['trade_date']} {r['action']} @ {r['price']}")


if __name__ == "__main__":
    main()
