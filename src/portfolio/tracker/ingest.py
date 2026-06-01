"""Ingest Excel trade ledger into parquet."""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path

import polars as pl

from src.portfolio.tracker.schema import TRADE_COLUMNS, validate_and_normalize
from src.utils.config import AppConfig, PortfolioTrackerConfig, load_app_config
from src.utils.logger import get_logger
from src.utils.paths import PROJECT_ROOT, ensure_dir

LOGGER = get_logger(__name__)


def _resolve_path(cfg: PortfolioTrackerConfig, relative: str) -> Path:
    path = Path(relative)
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    return path


def resolve_excel_path(app: AppConfig) -> tuple[Path, str]:
    """Return (excel path, label: primary|fallback)."""
    cfg = app.tracker
    input_dir = _resolve_path(cfg, cfg.input_dir)
    primary = input_dir / cfg.excel_primary
    if primary.is_file():
        return primary, "primary"
    fallback = input_dir / cfg.excel_fallback
    if fallback.is_file():
        return fallback, "fallback"
    raise FileNotFoundError(
        f"No trade Excel found. Expected {primary} or {fallback}. "
        "Copy dummy_portfolio.xlsx to input_trades.xlsx."
    )


def read_trades_excel(path: Path, *, sheet_name: str) -> pl.DataFrame:
    """Read trades sheet via openpyxl → Polars."""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise ValueError(f"Sheet {sheet_name!r} not found in {path.name}")
    ws = wb[sheet_name]
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    col_index = {name: headers.index(name) for name in TRADE_COLUMNS if name in headers}
    missing = [c for c in TRADE_COLUMNS if c not in col_index]
    if missing:
        wb.close()
        raise ValueError(f"Missing columns in {path.name}: {missing}")

    records: list[dict[str, object]] = []
    for row in rows_iter:
        if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        records.append(
            {col: row[col_index[col]] if col_index[col] < len(row) else None for col in TRADE_COLUMNS}
        )
    wb.close()
    return pl.DataFrame(records)


def ingest_trades(
    app: AppConfig | None = None,
    *,
    force: bool = False,
) -> Path:
    """Read Excel, validate, write trades parquet and metadata."""
    app = app or load_app_config()
    cfg = app.tracker
    excel_path, source_label = resolve_excel_path(app)
    out_path = _resolve_path(cfg, cfg.trades_parquet)
    meta_path = _resolve_path(cfg, cfg.metadata_json)

    if not force and out_path.is_file() and out_path.stat().st_mtime >= excel_path.stat().st_mtime:
        LOGGER.info("Trades parquet up to date: %s", out_path)
        return out_path

    raw = read_trades_excel(excel_path, sheet_name=cfg.sheet_name)
    trades = validate_and_normalize(raw)

    ensure_dir(out_path.parent)
    trades.write_parquet(out_path)

    metadata = {
        "source_excel": str(excel_path.relative_to(PROJECT_ROOT)),
        "source_label": source_label,
        "row_count": trades.height,
        "ingested_at": datetime.now(timezone.utc).isoformat(),
        "output_parquet": str(out_path.relative_to(PROJECT_ROOT)),
    }
    meta_path.write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    LOGGER.info(
        "Ingested %s trades from %s -> %s",
        trades.height,
        excel_path.name,
        out_path,
    )
    return out_path


def main() -> None:
    path = ingest_trades(force=True)
    print(path)


if __name__ == "__main__":
    main()
